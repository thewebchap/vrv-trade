"""Permanent trade storage: one folder per trade plus a master Excel register.

    <TRADE_STORAGE_PATH>/                     default: trade/
        trade_register.xlsx                   one row per trade (operational index)
        TRD-2026-000001/
            source/                           the trade's PDFs, as received
            result.json  report.md  report.xlsx  email_response.txt     (same writer as local mode)
            trade_metadata.json

Trade references are TRD-<year>-<6-digit sequence>. They are allocated under an exclusive file lock and
claimed by creating the folder with os.mkdir (atomic: it fails if the folder already exists), so two
processes can never receive the same reference.
"""
from __future__ import annotations

import fcntl
import json
import os
import re
import tempfile
from contextlib import contextmanager
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from .config import trade_storage_path

REF_RE = re.compile(r"^TRD-(\d{4})-(\d{6})$")
REGISTER = "trade_register.xlsx"
COLUMNS = ["Trade Ref", "Source", "Processing Status", "Received Date", "Processed Date", "Sender", "Email Subject",
           "Purchase Contract", "Purchase Invoice", "Bill of Lading", "Sales Contract", "Sales Invoice",
           "Other Documents", "PDF Count", "Pass Count", "Pass With Note Count", "Review Count", "Mismatch Count",
           "Fail Count", "Overall Status", "Payment Term", "Reply Status", "Trade Folder", "Notes"]
DOC_COLUMNS = {"PURCHASE_CONTRACT": "Purchase Contract", "PURCHASE_INVOICE": "Purchase Invoice",
               "BILL_OF_LADING": "Bill of Lading", "SALES_CONTRACT": "Sales Contract", "SALES_INVOICE": "Sales Invoice"}


class TradeStore:
    def __init__(self, root: str | None = None):
        self.root = os.path.abspath(root or trade_storage_path())
        os.makedirs(self.root, exist_ok=True)

    # ------------------------------------------------------------------ locking
    @contextmanager
    def _lock(self, name: str):
        with open(os.path.join(self.root, f".{name}.lock"), "w") as fh:
            fcntl.flock(fh, fcntl.LOCK_EX)
            try:
                yield
            finally:
                fcntl.flock(fh, fcntl.LOCK_UN)

    # ------------------------------------------------------------------ trade folders
    def allocate(self, year: int | None = None) -> str:
        """Next free TRD-<year>-NNNNNN; its folder (and source/) is created as part of allocation."""
        year = year or datetime.now().year
        with self._lock("allocate"):
            used = [int(m.group(2)) for d in os.listdir(self.root) if (m := REF_RE.match(d)) and int(m.group(1)) == year]
            n = max(used, default=0) + 1
            while True:
                ref = f"TRD-{year}-{n:06d}"
                try:
                    os.mkdir(self.folder(ref))
                    break
                except FileExistsError:
                    n += 1
            os.makedirs(self.source_folder(ref), exist_ok=True)
            return ref

    def folder(self, ref: str) -> str:
        if not REF_RE.match(ref):
            raise ValueError(f"not a trade reference: {ref!r}")
        return os.path.join(self.root, ref)

    def source_folder(self, ref: str) -> str:
        return os.path.join(self.folder(ref), "source")

    def save_source_pdf(self, ref: str, filename: str, content: bytes) -> str:
        """Write a PDF into source/ under a sanitised, non-colliding name. Returns the stored path."""
        folder = self.source_folder(ref)
        os.makedirs(folder, exist_ok=True)
        stem, _ = os.path.splitext(safe_filename(filename))
        candidate, i = f"{stem}.pdf", 1
        while True:
            path = os.path.join(folder, candidate)
            try:
                with open(path, "xb") as f:      # 'x': never overwrite an existing file
                    f.write(content)
                return path
            except FileExistsError:
                i += 1
                candidate = f"{stem}_{i}.pdf"

    def write_metadata(self, ref: str, metadata: dict) -> str:
        path = os.path.join(self.folder(ref), "trade_metadata.json")
        _atomic_write_text(path, json.dumps(metadata, indent=2, default=str))
        return path

    def read_metadata(self, ref: str) -> dict:
        path = os.path.join(self.folder(ref), "trade_metadata.json")
        if not os.path.exists(path):
            return {}
        with open(path) as f:
            return json.load(f)

    # ------------------------------------------------------------------ master register
    @property
    def register_path(self) -> str:
        return os.path.join(self.root, REGISTER)

    def upsert_register(self, row: dict) -> None:
        """Insert or update the row for row['Trade Ref']; existing rows are preserved. Saved atomically."""
        with self._lock("register"):
            if os.path.exists(self.register_path):
                wb = load_workbook(self.register_path)
                ws = wb.active
                headers = [c.value for c in ws[1]]
                for col in COLUMNS:                         # tolerate registers created before a column existed
                    if col not in headers:
                        ws.cell(1, len(headers) + 1, col).font = Font(bold=True)
                        headers.append(col)
            else:
                wb = Workbook()
                ws = wb.active
                ws.title = "Trades"
                ws.append(COLUMNS)
                for c in ws[1]:
                    c.font = Font(bold=True)
                ws.freeze_panes = "A2"
                headers = list(COLUMNS)
            ref_col = headers.index("Trade Ref") + 1
            target = next((r for r in range(2, ws.max_row + 1) if ws.cell(r, ref_col).value == row["Trade Ref"]),
                          ws.max_row + 1)
            for key, value in row.items():
                if key in headers:
                    ws.cell(target, headers.index(key) + 1, value)
            fd, tmp = tempfile.mkstemp(dir=self.root, suffix=".xlsx.tmp")
            os.close(fd)
            try:
                wb.save(tmp)
                os.replace(tmp, self.register_path)
            finally:
                if os.path.exists(tmp):
                    os.remove(tmp)

    def register_rows(self) -> list[dict]:
        if not os.path.exists(self.register_path):
            return []
        ws = load_workbook(self.register_path, read_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        return [dict(zip(rows[0], r)) for r in rows[1:]] if rows else []


def safe_filename(name: str | None, default: str = "document") -> str:
    """A plain file name: no directories, no traversal, no control or shell-special characters."""
    base = os.path.basename((name or "").replace("\\", "/"))
    base = re.sub(r"[^A-Za-z0-9._ -]", "_", base).strip(" .")
    base = re.sub(r"\s+", " ", base)[:120]
    return base or default


def _atomic_write_text(path: str, text: str) -> None:
    fd, tmp = tempfile.mkstemp(dir=os.path.dirname(path), suffix=".tmp")
    with os.fdopen(fd, "w") as f:
        f.write(text)
    os.replace(tmp, path)


def register_row(ref: str, res: dict | None, metadata: dict, folder: str) -> dict:
    """The register row for a trade, from its result (None when processing failed) and metadata."""
    docs = (res or {}).get("documents", [])
    by_type = {}
    for d in docs:
        by_type.setdefault(d["doc_type"], []).append(d["filename"])
    counts = ((res or {}).get("summary") or {}).get("checks") or {}
    other = [f for t, fs in by_type.items() if t not in DOC_COLUMNS for f in fs]
    return {
        "Trade Ref": ref, "Source": metadata.get("source"), "Processing Status": metadata.get("processing_status"),
        "Received Date": metadata.get("received_at"), "Processed Date": metadata.get("processed_at"),
        "Sender": metadata.get("sender"), "Email Subject": metadata.get("subject"),
        **{col: ", ".join(by_type.get(t, [])) or "missing" for t, col in DOC_COLUMNS.items()},
        "Other Documents": ", ".join(other) or None,
        "PDF Count": metadata.get("pdf_count"),
        "Pass Count": counts.get("PASS"), "Pass With Note Count": counts.get("PASS_WITH_NOTE"),
        "Review Count": counts.get("REVIEW"), "Mismatch Count": counts.get("MISMATCH"), "Fail Count": counts.get("FAIL"),
        "Overall Status": metadata.get("overall_status"),
        "Payment Term": ((res or {}).get("payment_term") or {}).get("rule_id"),
        "Reply Status": metadata.get("reply_status"), "Trade Folder": folder, "Notes": metadata.get("notes"),
    }
